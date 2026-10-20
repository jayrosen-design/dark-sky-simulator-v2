#!/usr/bin/env python3
"""Reconstruct the simple 2012 vs 2024 screenshot analysis from EPAC slide 5.

IMPORTANT: THIS IS A NEW RECONSTRUCTION, NOT THE LOST 2025 SCRIPT.

Inputs: the actual PNG screenshots embedded in the January 13, 2026
"Preserving Natural Skies in Alachua County" EPAC presentation.

Method: extract screenshots, register the 2024 image to the 2012 map,
select a fixed rectangular scene, convert RGB to grayscale, measure mean
pixel intensity and the fraction of pixels brighter than a fixed threshold.

This is screenshot-image analysis, NOT calibrated VIIRS radiometry.
The scene crop extends well beyond Alachua County, the threshold was chosen
post hoc in an attempt to reconstruct previously published values, and the
map includes yellow road lines, labels and other cartographic overlays.
Do NOT cite the results as a validated countywide radiance change.

Requirements:
    pip install python-pptx pillow opencv-python numpy matplotlib torch

Usage:
    python alachua_screenshot_reconstruction.py --pptx "EPAC.pptx" \
        --output-dir screenshot_results

The defaults (central scene [150, 50, 1430, 840] and threshold 38/255)
closely recreate the historical 19% and 16% -> 21% screenshot statistics.
Neither the historical crop nor its threshold has been recovered.
"""

from __future__ import annotations
import argparse
import csv
import io
from pathlib import Path
import cv2
import numpy as np
from PIL import Image
from pptx import Presentation
import torch
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

# Inclusive on left/top, exclusive on right/bottom; pixel coordinates of
# the full-resolution slide screenshots. Not a GIS county polygon!
DEFAULT_CROP = (150, 50, 1430, 840)
DEFAULT_THRESHOLD = 38  # gray intensity on scale 0..255, selected post hoc
REGIONS = {
    "Central screenshot crop (historical approximate)": DEFAULT_CROP,
    "Nearly full visible map (sensitivity test)": (10, 36, 1690, 1020),
    "Gainesville-area sample (not county boundary)": (740, 321, 890, 452),
    "Alachua-area sample (not city boundary)": (686, 276, 787, 343),
    "Paynes Prairie vicinity sample": (779, 401, 857, 469),
}


def extract_slide5_images(pptx_path: Path, out_dir: Path) -> tuple[Path, Path]:
    """Get the original, full-resolution images (top=2012, bottom=2024)."""
    presentation = Presentation(str(pptx_path))
    if len(presentation.slides) < 5:
        raise ValueError("Expected the EPAC presentation, with satellite maps on slide 5")
    images = [s for s in presentation.slides[4].shapes if s.shape_type == 13]
    images.sort(key=lambda s: s.top)
    if len(images) != 2:
        raise ValueError(f"Expected two pictures on slide 5, got {len(images)}")
    output = []
    for year, shape in zip((2012, 2024), images):
        image = Image.open(io.BytesIO(shape.image.blob)).convert("RGB")
        path = out_dir / f"EPAC_slide5_{year}_original.png"
        image.save(path)
        output.append(path)
    return tuple(output)


def register_2024_to_2012(img12: np.ndarray, img24: np.ndarray, radius=6):
    """Search small integer translations for best scene registration.

    Ref: 2012 screenshot. 2024 has a slight map pan (~3 px). We exclude
    the map controls/attribution during matching. No scaling or warping.
    """
    a = cv2.cvtColor(img12, cv2.COLOR_RGB2GRAY)
    b = cv2.cvtColor(img24, cv2.COLOR_RGB2GRAY)
    left, top, right, bottom = (200, 90, 1400, 810)
    ref = a[top:bottom, left:right].astype(np.int16)
    best = (float("inf"), 0, 0)
    for dy in range(-radius, radius + 1):
        for dx in range(-radius, radius + 1):
            moving = b[top - dy:bottom - dy, left - dx:right - dx]
            score = np.abs(ref - moving.astype(np.int16)).mean()
            if score < best[0]:
                best = (score, dx, dy)
    _, dx, dy = best
    warped = cv2.warpAffine(
        img24, np.float32([[1, 0, dx], [0, 1, dy]]),
        (img12.shape[1], img12.shape[0]),
        flags=cv2.INTER_NEAREST,
        borderMode=cv2.BORDER_CONSTANT,
    )
    return warped, {"dx": dx, "dy": dy, "match_mean_abs_difference": float(best[0])}


def measure(img12: np.ndarray, img24: np.ndarray, crop, threshold: int):
    """Actual pixel math uses PyTorch CPU tensors; no trained AI required."""
    x0, y0, x1, y1 = crop
    gray12 = cv2.cvtColor(img12[y0:y1, x0:x1], cv2.COLOR_RGB2GRAY)
    gray24 = cv2.cvtColor(img24[y0:y1, x0:x1], cv2.COLOR_RGB2GRAY)
    if gray12.shape != gray24.shape or not gray12.size:
        raise ValueError("Invalid crop or incompatible image sizes")
    a = torch.from_numpy(gray12.astype(np.float32))
    b = torch.from_numpy(gray24.astype(np.float32))
    mean12 = float(a.mean().item())
    mean24 = float(b.mean().item())
    if mean12 <= 0:
        raise ValueError("Zero mean in 2012 image")
    return {
        "roi": str(tuple(crop)),
        "threshold_0_255": threshold,
        "pixel_count": int(a.numel()),
        "mean_gray_2012": mean12,
        "mean_gray_2024": mean24,
        "mean_gray_change_pct": (mean24 / mean12 - 1) * 100,
        "lit_pixels_2012_pct": float((a > threshold).float().mean().item() * 100),
        "lit_pixels_2024_pct": float((b > threshold).float().mean().item() * 100),
        "lit_change_percentage_points": float(
            ((b > threshold).float().mean() - (a > threshold).float().mean()).item() * 100
        ),
    }


def visualization(img12, img24, crop, threshold, stats, output_file):
    x0, y0, x1, y1 = crop
    a = img12[y0:y1, x0:x1]
    b = img24[y0:y1, x0:x1]
    ga = cv2.cvtColor(a, cv2.COLOR_RGB2GRAY)
    gb = cv2.cvtColor(b, cv2.COLOR_RGB2GRAY)
    increase = np.maximum(gb.astype(np.int16) - ga.astype(np.int16), 0).astype(np.uint8)
    # Display enhancement is ONLY for plotting and not used in calculations.
    change_view = cv2.GaussianBlur(increase, (0, 0), sigmaX=3)
    fig, axes = plt.subplots(3, 1, figsize=(15, 14), constrained_layout=True)
    axes[0].imshow(a)
    axes[0].set_title(f"2012 source screenshot / central map crop (mean grayscale {stats['mean_gray_2012']:.2f})")
    axes[1].imshow(b)
    axes[1].set_title(f"2024 source screenshot / registered (mean grayscale {stats['mean_gray_2024']:.2f})")
    axes[2].imshow(change_view, cmap="gray")
    axes[2].set_title("Positive grayscale changes (illustration; background and map overlay artifacts remain)")
    for ax in axes:
        ax.axis("off")
    fig.suptitle(
        f"Screenshot reconstruction — mean +{stats['mean_gray_change_pct']:.1f}%  |  "
        f"pixels > {threshold}: {stats['lit_pixels_2012_pct']:.1f}% → {stats['lit_pixels_2024_pct']:.1f}%\n"
        "Crop and threshold selected after seeing historical figures; not an independently validated county trend",
        fontsize=13,
    )
    fig.savefig(output_file, dpi=145)
    plt.close(fig)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--pptx", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, default=Path("screenshot_results"))
    parser.add_argument("--threshold", type=int, default=DEFAULT_THRESHOLD)
    parser.add_argument("--crop", nargs=4, metavar=("LEFT", "TOP", "RIGHT", "BOTTOM"),
                        type=int, default=DEFAULT_CROP)
    args = parser.parse_args()
    if not 0 <= args.threshold <= 255:
        parser.error("Threshold must be within 0..255")
    args.output_dir.mkdir(parents=True, exist_ok=True)
    p12, p24 = extract_slide5_images(args.pptx, args.output_dir)
    a = np.array(Image.open(p12).convert("RGB"))
    b_raw = np.array(Image.open(p24).convert("RGB"))
    b, alignment = register_2024_to_2012(a, b_raw)
    Image.fromarray(b).save(args.output_dir / "EPAC_2024_registered.png")

    all_regions = dict(REGIONS)
    all_regions["Requested crop"] = tuple(args.crop)
    rows = []
    for name, crop in all_regions.items():
        row = {"region": name, **measure(a, b, crop, args.threshold)}
        rows.append(row)
    csv_path = args.output_dir / "screenshot_stats.csv"
    with csv_path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
    stats = rows[-1]
    visualization(a, b, tuple(args.crop), args.threshold, stats,
                  args.output_dir / "screenshot_comparison.png")
    print(f"2012: {p12}")
    print(f"2024: {p24}")
    print(f"Registration (2024 -> 2012): {alignment}")
    print(f"Threshold: {args.threshold}/255, selected post hoc")
    for row in rows:
        print(f"{row['region']}: brightness {row['mean_gray_change_pct']:+.2f}%; "
              f"lit {row['lit_pixels_2012_pct']:.2f}% -> "
              f"{row['lit_pixels_2024_pct']:.2f}%")
    print(f"Full CSV: {csv_path}")

if __name__ == "__main__":
    main()
