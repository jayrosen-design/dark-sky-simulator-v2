#!/usr/bin/env python3
"""Measure the EPAC slide-5 night-lights screenshots inside real boundaries.

Refactor of alachua_screenshot_reconstruction.py (October 2026). The first reconstruction reproduced the published
+19% and 16% -> 21% only with a 212 x 131 km rectangle (Perry to the St. Johns River, Lake City to Belleview) and a
threshold chosen after the fact. This version measures the places the published figures name:

    Alachua County            Census TIGER 2020 county boundary
    Gainesville               2020 city limits (Census incorporated place)
    City of Alachua           2020 city limits (Census incorporated place)
    Paynes Prairie vicinity   Paynes Prairie Preserve State Park (FNAI), alone and with 2, 5 and 10 km rings

and, because the original may have used rectangles, the whole map view, each place's bounding box and the rectangles
the 2025 Dark Sky Simulator drew (rectangle_stats.csv). Whole map + the three place boxes reproduce the published set.

How: the screenshots are Google Maps tiles, i.e. Web Mercator, so pixel x is linear in longitude and pixel y in
Mercator northing. Both are fitted to nine town labels on the 2012 image (median residual about 0.3 km); the 2024
image is registered onto the 2012 image by the same integer-shift search as before. Every pixel center is then tested
against the boundaries in boundaries.geojson.

What the numbers are: change in screenshot grayscale (ITU-R 601 luma of a rendered color ramp, with road lines and
labels on top), not satellite radiance. Use them to understand the published figures, not as a measured trend.

Usage:
    python alachua_screenshot_regions.py --images EPAC_slide5_2012_original.png EPAC_slide5_2024_original.png
    python alachua_screenshot_regions.py --pptx "EPAC.pptx"            # extracts slide 5 (needs python-pptx)
    add --variants to score 24 pre-declared metric variants against the published numbers.

Requirements: numpy. Optional: Pillow (faster PNG reading), matplotlib (outline figure), python-pptx (--pptx).
"""
from __future__ import annotations

import argparse
import csv
import io
import json
import math
import struct
import zlib
from pathlib import Path

import numpy as np

HERE = Path(__file__).parent
# Figures published on EPAC slide 5 (January 13, 2026); compared against, never fitted to.
PUBLISHED_PCT = {"alachua_county": 19, "gainesville": 6, "alachua_city": 10, "paynes_prairie_5km": 15}
PUBLISHED_LIT = (16, 21)
# Town label centers on the 2012 slide image (original pixels) and the towns' coordinates.
CONTROL_POINTS = {
    "Gainesville": (984, 458, 29.6516, -82.3248), "Ocala": (1096, 771, 29.1872, -82.1401),
    "Lake City": (801, 95, 30.1897, -82.6393), "Palatka": (1390, 461, 29.6486, -81.6376),
    "Cedar Key": (569, 802, 29.1386, -83.0351), "Perry": (251, 148, 30.1174, -83.5821),
    "Starke": (1111, 260, 29.9441, -82.1098), "Chiefland": (672, 568, 29.4749, -82.8598),
    "Williston": (914, 636, 29.3874, -82.4468),
}
ORIGINAL_CROP = (150, 50, 1430, 840)   # the first reconstruction's rectangle, kept for comparison
# Rectangles the 2025 Dark Sky Simulator drew on its map (south, north, west, east), github.com/jayrosen-design/Dark-Sky-Simulator
APP_2025_RECTANGLES = {"Gainesville Downtown (2025 app)": (29.64, 29.67, -82.35, -82.32),
                       "Suburban Gainesville (2025 app)": (29.60, 29.70, -82.40, -82.25),
                       "Paynes Prairie (2025 app)": (29.55, 29.62, -82.35, -82.25)}


# ---------------------------------------------------------------- images

def read_rgb(path: Path) -> np.ndarray:
    """RGB array of a PNG: Pillow when installed, else a small reader for 8-bit, non-interlaced PNGs."""
    try:
        from PIL import Image
        return np.asarray(Image.open(path).convert("RGB"), dtype=float)
    except ImportError:
        pass
    data = path.read_bytes()
    pos, idat = 8, b""
    while pos < len(data):
        n, typ = struct.unpack(">I4s", data[pos:pos + 8])
        body = data[pos + 8:pos + 8 + n]
        if typ == b"IHDR":
            w, h, depth, ctype, _, _, interlace = struct.unpack(">IIBBBBB", body)
            if depth != 8 or interlace or ctype not in (2, 6):
                raise ValueError(f"{path}: install Pillow to read this PNG")
            ch = 3 if ctype == 2 else 4
        elif typ == b"IDAT":
            idat += body
        pos += 12 + n
    raw = np.frombuffer(zlib.decompress(idat), dtype=np.uint8).reshape(h, 1 + w * ch)
    out = np.zeros((h, w * ch), dtype=np.int32)
    for r in range(h):
        f, line = raw[r, 0], raw[r, 1:].astype(np.int32)
        prev = out[r - 1] if r else np.zeros(w * ch, dtype=np.int32)
        if f == 0:
            out[r] = line
        elif f == 2:
            out[r] = (line + prev) & 255
        else:
            cur = out[r]
            for i in range(w * ch):
                left = cur[i - ch] if i >= ch else 0
                up, ul = prev[i], (prev[i - ch] if i >= ch else 0)
                if f == 1:
                    pred = left
                elif f == 3:
                    pred = (left + up) // 2
                else:
                    pa, pb, pc = abs(up - ul), abs(left - ul), abs(left + up - 2 * ul)
                    pred = left if pa <= pb and pa <= pc else up if pb <= pc else ul
                cur[i] = (line[i] + pred) & 255
    return out.reshape(h, w, ch)[..., :3].astype(float)


def luma(rgb: np.ndarray) -> np.ndarray:
    """Grayscale as cv2.COLOR_RGB2GRAY and PIL convert('L') compute it."""
    return np.floor(rgb[..., 0] * 0.299 + rgb[..., 1] * 0.587 + rgb[..., 2] * 0.114 + 0.5)


def extract_slide5(pptx: Path, out_dir: Path) -> tuple[Path, Path]:
    from pptx import Presentation
    slide = Presentation(str(pptx)).slides[4]
    pics = sorted((s for s in slide.shapes if s.shape_type == 13), key=lambda s: s.top)
    if len(pics) != 2:
        raise ValueError(f"expected two pictures on slide 5, found {len(pics)}")
    paths = []
    for year, pic in zip((2012, 2024), pics):
        p = out_dir / f"EPAC_slide5_{year}_original.png"
        p.write_bytes(pic.image.blob)
        paths.append(p)
    return paths[0], paths[1]


def register(rgb12: np.ndarray, rgb24: np.ndarray, radius: int = 6):
    """Best integer shift of the 2024 image onto the 2012 image (mean absolute gray difference, map controls excluded)."""
    a, b = luma(rgb12), luma(rgb24)
    l, t, r, btm = 200, 90, 1400, 810
    ref = a[t:btm, l:r]
    best = min(((np.abs(ref - b[t - dy:btm - dy, l - dx:r - dx]).mean(), dx, dy)
                for dy in range(-radius, radius + 1) for dx in range(-radius, radius + 1)))
    _, dx, dy = best
    # Output on the 2012 image's grid (the two screenshots differ in size): out(y, x) = 2024(y - dy, x - dx), else 0.
    src_y, src_x = np.arange(rgb12.shape[0]) - dy, np.arange(rgb12.shape[1]) - dx
    ry, rx = (src_y >= 0) & (src_y < rgb24.shape[0]), (src_x >= 0) & (src_x < rgb24.shape[1])
    out = np.zeros_like(rgb12)
    out[np.ix_(ry, rx)] = rgb24[np.ix_(src_y[ry], src_x[rx])]
    return out, {"dx": int(dx), "dy": int(dy), "mean_abs_difference": float(best[0])}


# ---------------------------------------------------------------- geography

def georeference(shape: tuple[int, int]):
    """Longitude and latitude of every pixel center, from a Web Mercator fit to the control points."""
    merc = lambda lat: np.log(np.tan(np.pi / 4 + np.radians(lat) / 2))
    cp = np.array(list(CONTROL_POINTS.values()), dtype=float)
    bx = np.linalg.lstsq(np.column_stack([np.ones(len(cp)), cp[:, 3]]), cp[:, 0], rcond=None)[0]
    by = np.linalg.lstsq(np.column_stack([np.ones(len(cp)), merc(cp[:, 2])]), cp[:, 1], rcond=None)[0]
    resid = np.hypot(bx[0] + bx[1] * cp[:, 3] - cp[:, 0], by[0] + by[1] * merc(cp[:, 2]) - cp[:, 1])
    km_px = 111.32 * math.cos(math.radians(29.6)) / bx[1]
    ys, xs = np.mgrid[0:shape[0], 0:shape[1]]
    lon = (xs + 0.5 - bx[0]) / bx[1]
    lat = np.degrees(2 * np.arctan(np.exp((ys + 0.5 - by[0]) / by[1])) - np.pi / 2)
    return lon, lat, {"median_residual_px": float(np.median(resid)), "km_per_px": float(km_px)}


def inside(geometry: dict, lon: np.ndarray, lat: np.ndarray) -> np.ndarray:
    """Even-odd point-in-polygon test of every pixel against a GeoJSON (Multi)Polygon, holes included."""
    polys = geometry["coordinates"] if geometry["type"] == "MultiPolygon" else [geometry["coordinates"]]
    rings = [np.asarray(r, dtype=float) for poly in polys for r in poly]
    allpts = np.vstack(rings)
    sel = (lon >= allpts[:, 0].min()) & (lon <= allpts[:, 0].max()) & (lat >= allpts[:, 1].min()) & (lat <= allpts[:, 1].max())
    x, y = lon[sel], lat[sel]
    odd = np.zeros(x.shape, bool)
    for ring in rings:
        for (x1, y1), (x2, y2) in zip(ring[:-1], ring[1:]):
            crosses = (y1 > y) != (y2 > y)
            with np.errstate(divide="ignore", invalid="ignore"):
                xc = x1 + (y - y1) * (x2 - x1) / (y2 - y1)
            odd ^= crosses & (x < xc)
    m = np.zeros(lon.shape, bool)
    m[sel] = odd
    return m


def bbox_mask(geometry: dict, lon: np.ndarray, lat: np.ndarray) -> np.ndarray:
    pts = np.vstack([np.asarray(r) for poly in (geometry["coordinates"] if geometry["type"] == "MultiPolygon"
                                                   else [geometry["coordinates"]]) for r in poly])
    return (lon >= pts[:, 0].min()) & (lon <= pts[:, 0].max()) & (lat >= pts[:, 1].min()) & (lat <= pts[:, 1].max())


# ---------------------------------------------------------------- statistics

def measure(g12: np.ndarray, g24: np.ndarray, m: np.ndarray, threshold: int) -> dict:
    a, b = g12[m], g24[m]
    return {"pixels": int(m.sum()), "mean_gray_2012": a.mean(), "mean_gray_2024": b.mean(),
            "change_pct": (b.mean() / a.mean() - 1) * 100,
            "lit_2012_pct": (a > threshold).mean() * 100, "lit_2024_pct": (b > threshold).mean() * 100}


CHANNELS = {"luma (cv2/PIL gray)": luma, "value, max(R,G,B)": lambda c: c.max(axis=-1),
            "mean of R,G,B": lambda c: c.mean(axis=-1),
            "yellow, (R+G)/2 - B": lambda c: np.maximum((c[..., 0] + c[..., 1]) / 2 - c[..., 2], 0)}
STATISTICS = {"mean": lambda v: v.mean(), "sum above 38": lambda v: v[v > 38].sum(),
              "lit pixel count (>38)": lambda v: float((v > 38).sum())}


def variants(rgb12, rgb24, outline: dict, box: dict) -> list[dict]:
    """24 pre-declared variants (4 channels x 2 region shapes x 3 statistics), each scored on all four published
    numbers at once (RMS error, percentage points). With 24 tries, a close match can happen by chance."""
    rows = []
    for cn, cf in CHANNELS.items():
        a, b = cf(rgb12), cf(rgb24)
        for shape_name, masks in (("outline", outline), ("bounding box", box)):
            for sn, sf in STATISTICS.items():
                pct = {k: (sf(b[masks[k]]) / sf(a[masks[k]]) - 1) * 100 for k in PUBLISHED_PCT}
                err = math.sqrt(np.mean([(pct[k] - PUBLISHED_PCT[k]) ** 2 for k in PUBLISHED_PCT]))
                rows.append({"rms_error_pts": err, "channel": cn, "region_shape": shape_name, "statistic": sn,
                             **{f"{k}_pct": v for k, v in pct.items()}})
    return sorted(rows, key=lambda r: r["rms_error_pts"])


def figure(rgb12, rgb24, masks: dict, out: Path):
    """2012 and 2024 around Alachua County with the measured outlines, if matplotlib is installed."""
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except ImportError:
        return False
    rows, cols = np.nonzero(masks["alachua_county"])
    y0, y1, x0, x1 = rows.min() - 40, rows.max() + 40, cols.min() - 40, cols.max() + 40
    fig, axes = plt.subplots(1, 2, figsize=(16, 8), constrained_layout=True)
    colors = {"alachua_county": "white", "gainesville": "cyan", "alachua_city": "magenta", "paynes_prairie_park": "lime",
              "paynes_prairie_5km": "lime"}
    for ax, img, year in zip(axes, (rgb12, rgb24), (2012, 2024)):
        ax.imshow(img[y0:y1, x0:x1].astype(np.uint8))
        for k, c in colors.items():
            ax.contour(masks[k][y0:y1, x0:x1], levels=[0.5], colors=c, linewidths=1.2,
                       linestyles="--" if k.endswith("km") else "-")
        ax.set_title(f"{year} screenshot, measured areas")
        ax.axis("off")
    fig.suptitle("White: Alachua County. Cyan: Gainesville. Magenta: City of Alachua. Green: Paynes Prairie (dashed: +5 km)")
    fig.savefig(out, dpi=110)
    plt.close(fig)
    return True


def write_csv(path: Path, rows: list[dict]):
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows([{k: round(v, 3) if isinstance(v, float) else v for k, v in r.items()} for r in rows])


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    src = ap.add_mutually_exclusive_group(required=True)
    src.add_argument("--images", nargs=2, type=Path, metavar=("PNG_2012", "PNG_2024"))
    src.add_argument("--pptx", type=Path)
    ap.add_argument("--boundaries", type=Path, default=HERE / "boundaries.geojson")
    ap.add_argument("--output-dir", type=Path, default=Path("region_results"))
    ap.add_argument("--threshold", type=int, default=38, help="gray level counted as lit (0-255)")
    ap.add_argument("--variants", action="store_true", help="also score 24 pre-declared metric variants")
    args = ap.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)

    p12, p24 = extract_slide5(args.pptx, args.output_dir) if args.pptx else args.images
    rgb12 = read_rgb(Path(p12))
    rgb24, shift = register(rgb12, read_rgb(Path(p24)))
    g12, g24 = luma(rgb12), luma(rgb24)
    lon, lat, geo = georeference(g12.shape)
    feats = {f["properties"]["id"]: f for f in json.loads(args.boundaries.read_text())["features"]}
    masks = {k: inside(f["geometry"], lon, lat) for k, f in feats.items()}
    masks["paynes_prairie_5km_outside_gainesville"] = masks["paynes_prairie_5km"] & ~masks["gainesville"]
    x0, y0, x1, y1 = ORIGINAL_CROP
    masks["original_rectangle"] = np.zeros(g12.shape, bool)
    masks["original_rectangle"][y0:y1, x0:x1] = True
    names = {k: f["properties"]["name"] for k, f in feats.items()}
    names["paynes_prairie_5km_outside_gainesville"] = "Paynes Prairie + 5 km, outside Gainesville"
    names["original_rectangle"] = "First reconstruction's rectangle (212 x 131 km)"

    print(f"Registration (2024 onto 2012): {shift}")
    print(f"Georeference: median residual {geo['median_residual_px']:.1f} px, {geo['km_per_px']:.2f} km per pixel")
    rows = []
    for k in ("alachua_county", "gainesville", "alachua_city", "paynes_prairie_park", "paynes_prairie_2km",
              "paynes_prairie_5km", "paynes_prairie_10km", "paynes_prairie_5km_outside_gainesville", "original_rectangle"):
        r = {"region": names[k], **measure(g12, g24, masks[k], args.threshold),
             "published_pct": PUBLISHED_PCT.get(k, PUBLISHED_PCT["alachua_county"] if k == "original_rectangle" else "")}
        rows.append(r)
        print(f"{r['region']:<48} {r['pixels']:>8} px  gray {r['mean_gray_2012']:6.2f} -> {r['mean_gray_2024']:6.2f}"
              f"  {r['change_pct']:+6.1f}%   published {r['published_pct']}")
    write_csv(args.output_dir / "region_stats.csv", rows)

    # Rectangles instead of boundaries: the whole map view, each place's bounding box, and the 2025 app's rectangles.
    rects = {"Whole map view (first reconstruction's crop)": (masks["original_rectangle"], "alachua_county"),
             "Alachua County, bounding box": (bbox_mask(feats["alachua_county"]["geometry"], lon, lat), "alachua_county"),
             "Gainesville, bounding box": (bbox_mask(feats["gainesville"]["geometry"], lon, lat), "gainesville"),
             "City of Alachua, bounding box": (bbox_mask(feats["alachua_city"]["geometry"], lon, lat), "alachua_city"),
             "Paynes Prairie Preserve, bounding box": (bbox_mask(feats["paynes_prairie_park"]["geometry"], lon, lat), "paynes_prairie_5km")}
    for name, (s, n, w, e) in APP_2025_RECTANGLES.items():
        rects[name] = ((lat >= s) & (lat <= n) & (lon >= w) & (lon <= e), "paynes_prairie_5km" if "Paynes" in name else "gainesville")
    print("\nRectangles:")
    rrows = []
    for name, (mk, pub) in rects.items():
        r = {"rectangle": name, **measure(g12, g24, mk, args.threshold), "published_pct": PUBLISHED_PCT[pub]}
        rrows.append(r)
        print(f"{name:<48} {r['pixels']:>8} px  gray {r['mean_gray_2012']:6.2f} -> {r['mean_gray_2024']:6.2f}"
              f"  {r['change_pct']:+6.1f}%   published {r['published_pct']}")
    write_csv(args.output_dir / "rectangle_stats.csv", rrows)
    best_set = [rrows[0], rrows[2], rrows[3], rrows[4]]   # whole map, Gainesville, City of Alachua, Paynes Prairie boxes
    err = math.sqrt(np.mean([(r["change_pct"] - r["published_pct"]) ** 2 for r in best_set]))
    print(f"Whole map + Gainesville, City of Alachua and Paynes Prairie boxes vs the four published numbers: RMS error {err:.1f} pts\n")

    m = masks["alachua_county"]
    sweep = [{"threshold": t, "lit_2012_pct": (g12[m] > t).mean() * 100, "lit_2024_pct": (g24[m] > t).mean() * 100}
             for t in range(20, 81, 5)]
    write_csv(args.output_dir / "county_threshold_sweep.csv", sweep)
    closest = min(sweep, key=lambda s: abs(s["lit_2012_pct"] - PUBLISHED_LIT[0]) + abs(s["lit_2024_pct"] - PUBLISHED_LIT[1]))
    print(f"County lit share: threshold {args.threshold} gives {(g12[m] > args.threshold).mean() * 100:.1f}% -> "
          f"{(g24[m] > args.threshold).mean() * 100:.1f}%; closest to the published {PUBLISHED_LIT[0]}% -> {PUBLISHED_LIT[1]}% "
          f"is threshold {closest['threshold']}: {closest['lit_2012_pct']:.1f}% -> {closest['lit_2024_pct']:.1f}%")

    if args.variants:
        boxes = {k: bbox_mask(feats[k]["geometry"], lon, lat) for k in PUBLISHED_PCT}
        v = variants(rgb12, rgb24, {k: masks[k] for k in PUBLISHED_PCT}, boxes)
        write_csv(args.output_dir / "metric_variants.csv", v)
        b = v[0]
        print(f"Best of 24 variants ({b['channel']}, {b['region_shape']}, {b['statistic']}): RMS error {b['rms_error_pts']:.1f} pts; "
              f"county {b['alachua_county_pct']:+.1f}, Gainesville {b['gainesville_pct']:+.1f}, "
              f"City of Alachua {b['alachua_city_pct']:+.1f}, Paynes Prairie {b['paynes_prairie_5km_pct']:+.1f}")
    if figure(rgb12, rgb24, masks, args.output_dir / "measured_areas.png"):
        print(f"Figure: {args.output_dir / 'measured_areas.png'}")
    print(f"CSV files in {args.output_dir}")


if __name__ == "__main__":
    main()
